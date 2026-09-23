"""
Tests for zimra_qpd.xlsx_export.

Two tiers:
  1. Structural/smoke tests (always run) - the workbook builds, has the
     right sheets, and pulled Inputs-sheet values match what was passed in.
     These don't need anything beyond openpyxl.
  2. Formula-correctness tests (skipped if LibreOffice's `soffice` isn't on
     PATH) - actually opens the generated workbook in a real spreadsheet
     engine, forces a recalculation, and diffs the Working Papers sheet's
     computed values against the real Python engine's own output for the
     SAME inputs (via worked_examples.run_company). This is the test that
     actually proves "all the formulas are correct", not just "a formula
     string got written somewhere".
"""

from __future__ import annotations

import dataclasses
import shutil
import subprocess
import sys

import pytest
from openpyxl import load_workbook

from zimra_qpd.worked_examples import KHAMI, UMGUZA, run_company
from zimra_qpd.xlsx_export import QuarterExportData, build_qpd_workbook

HAS_SOFFICE = shutil.which("soffice") is not None


def _quarters_from_company(company) -> tuple[list[QuarterExportData], list]:
    runs = run_company(company)
    quarters = []
    for run in runs:
        quarters.append(
            QuarterExportData(
                quarter=run.quarter,
                quarter_label=f"Q{run.quarter}",
                due_date=run.result.due_date,
                has_calculation=True,
                input=dataclasses.asdict(run.engine_input),
                result=dataclasses.asdict(run.result),
                actual_usd_paid=run.paid_usd,
                actual_zig_paid=run.paid_zig,
            )
        )
    return quarters, runs


# --- Tier 1: structural / smoke tests ---------------------------------


def test_workbook_has_expected_sheets():
    quarters, _ = _quarters_from_company(KHAMI)
    wb = build_qpd_workbook(KHAMI.name, KHAMI.tax_year, quarters)
    assert wb.sheetnames == ["Read Me", "Inputs", "Working Papers"]


def test_rejects_wrong_quarter_count():
    quarters, _ = _quarters_from_company(KHAMI)
    with pytest.raises(ValueError):
        build_qpd_workbook(KHAMI.name, KHAMI.tax_year, quarters[:3])


def test_rejects_out_of_order_quarters():
    quarters, _ = _quarters_from_company(KHAMI)
    shuffled = [quarters[1], quarters[0], quarters[2], quarters[3]]
    with pytest.raises(ValueError):
        build_qpd_workbook(KHAMI.name, KHAMI.tax_year, shuffled)


def test_missing_quarter_is_marked_not_calculated():
    quarters, _ = _quarters_from_company(KHAMI)
    quarters[3] = QuarterExportData(
        quarter=4, quarter_label="Q4", due_date="20 December", has_calculation=False
    )
    wb = build_qpd_workbook(KHAMI.name, KHAMI.tax_year, quarters)
    inputs_ws = wb["Inputs"]
    # Header row 1, column F is Q4 - should mention "not yet calculated".
    assert "not yet calculated" in str(inputs_ws["F1"].value).lower()


def test_inputs_sheet_pulls_the_confirmed_annual_sales_figure():
    quarters, runs = _quarters_from_company(KHAMI)
    wb = build_qpd_workbook(KHAMI.name, KHAMI.tax_year, quarters)
    inputs_ws = wb["Inputs"]

    # Find the "ANNUAL SALES ESTIMATE - USD" row and check Q1's value (col C)
    # matches what the engine was actually given for Q1.
    row = next(
        r for r in range(1, inputs_ws.max_row + 1)
        if inputs_ws[f"B{r}"].value and "ANNUAL SALES ESTIMATE" in str(inputs_ws[f"B{r}"].value)
        and "USD" in str(inputs_ws[f"B{r}"].value)
    )
    assert inputs_ws[f"C{row}"].value == pytest.approx(runs[0].engine_input.usd_sales)


# --- Tier 2: LibreOffice recalculation vs. the real engine -------------


@pytest.mark.skipif(not HAS_SOFFICE, reason="LibreOffice (soffice) not available on PATH")
@pytest.mark.parametrize("company", [KHAMI, UMGUZA], ids=lambda c: c.key)
def test_workbook_formulas_match_the_engine(tmp_path, company):
    quarters, runs = _quarters_from_company(company)
    wb = build_qpd_workbook(company.name, company.tax_year, quarters)

    src_path = tmp_path / f"{company.key}.xlsx"
    wb.save(src_path)

    recalced_dir = tmp_path / "recalced"
    recalced_dir.mkdir()
    subprocess.run(
        [
            "soffice", "--headless", "--calc",
            "--convert-to", "xlsx:Calc MS Excel 2007 XML",
            "--outdir", str(recalced_dir), str(src_path),
        ],
        check=True,
        capture_output=True,
        timeout=120,
    )

    recalced = load_workbook(recalced_dir / f"{company.key}.xlsx", data_only=True)
    ws = recalced["Working Papers"]

    label_row = {
        str(ws[f"B{r}"].value): r
        for r in range(1, ws.max_row + 1)
        if ws[f"B{r}"].value
    }
    cols = ["C", "D", "E", "F"]

    checks = [
        ("TAXABLE PROFIT \u2014 USD", lambda r: r.result.taxable_profit_usd),
        ("TAXABLE PROFIT \u2014 ZiG", lambda r: r.result.taxable_profit_zig),
        ("TOTAL TAX FOR THE YEAR AT THIS ESTIMATE \u2014 USD", lambda r: r.result.total_tax_usd),
        ("TOTAL TAX FOR THE YEAR AT THIS ESTIMATE \u2014 ZiG", lambda r: r.result.total_tax_zig),
        ("NET PAYABLE ON THIS QPD \u2014 USD", lambda r: r.result.net_payable_usd),
        ("NET PAYABLE ON THIS QPD \u2014 ZiG", lambda r: r.result.net_payable_zig),
    ]

    for label, getter in checks:
        row = label_row[label]
        for col, run in zip(cols, runs):
            expected = round(getter(run), 2)
            actual = ws[f"{col}{row}"].value
            assert actual is not None, f"{company.key} Q{run.quarter} {label}: formula produced no value"
            assert round(actual, 2) == pytest.approx(expected, abs=0.005), (
                f"{company.key} Q{run.quarter} {label}: expected {expected}, "
                f"workbook formula recalculated to {round(actual, 2)}"
            )
